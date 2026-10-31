"""Ekspor laporan Harian, Mingguan, Bulanan, dan Riwayat Siswa ke CSV dan XLSX."""
import csv
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import timeutil
from .attendance import STATUS_ORDER, STATUSES

_DANGEROUS = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value):
    """Cegah CSV/Excel formula injection pada isian yang berasal dari pengguna (nama, keterangan)."""
    if isinstance(value, str) and value.startswith(_DANGEROUS):
        return "'" + value
    return value


def _csv_bytes(header, rows):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    for row in rows:
        writer.writerow([safe_cell(c) for c in row])
    return buf.getvalue().encode("utf-8-sig")  # BOM agar Excel membaca UTF-8 dengan benar


def _xlsx_bytes(sheets):
    """sheets: daftar (judul, header, baris)."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, header, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="12306A")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in rows:
            ws.append([safe_cell(c) for c in row])
        for idx, col in enumerate(header, start=1):
            longest = max([len(str(col))] + [len(str(r[idx - 1])) for r in rows if idx - 1 < len(r)] or [8])
            ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 6), 40)
        ws.freeze_panes = "C2" if len(header) > 3 else "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME = "text/csv; charset=utf-8"


def _status(code):
    return STATUSES.get(code, "Belum presensi")


def _source_label(source):
    if not source:
        return ""
    return "Siswa" if source == "siswa" else "Admin kelas"


def daily_export(rows, d, fmt):
    header = ["No. Absen", "Nama Lengkap", "Status", "Waktu Presensi", "Sumber", "Keterangan"]
    data = [[r["roll_no"] or "", r["full_name"], _status(r["status"]),
             timeutil.fmt_time(r["submitted_at"]) if r["submitted_at"] else "",
             _source_label(r["source"]), r["note"] or ""] for r in rows]
    name = f"presensi-harian-{d}"
    if fmt == "xlsx":
        return _xlsx_bytes([(f"Harian {d}", header, data)]), XLSX_MIME, name + ".xlsx"
    return _csv_bytes(header, data), CSV_MIME, name + ".csv"


def weekly_export(report, fmt):
    start, end = report["start"], report["end"]
    week_cols = [
        f"{timeutil.DAYS_ID[date.fromisoformat(d).weekday()][:3]} {d[8:10]}/{d[5:7]}"
        for d in report["week_days"]
    ]
    header = ["No. Absen", "Nama Lengkap"] + week_cols + ["Hadir", "Izin", "Sakit", "Alpa", "Belum tercatat", "% Hadir"]
    data = [
        [
            r["roll_no"] or "",
            r["full_name"],
            *[r["cells"].get(d, "") for d in report["week_days"]],
            *[r["counts"][k] for k in STATUS_ORDER],
            r["belum"],
            "" if r["pct"] is None else r["pct"],
        ]
        for r in report["rows"]
    ]
    name = f"rekap-mingguan-{start.isoformat()}-sd-{end.isoformat()}"
    if fmt == "xlsx":
        summary_header = ["No. Absen", "Nama Lengkap", "Hadir", "Izin", "Sakit", "Alpa", "Belum tercatat", "% Hadir"]
        summary_data = [
            [
                r["roll_no"] or "",
                r["full_name"],
                *[r["counts"][k] for k in STATUS_ORDER],
                r["belum"],
                "" if r["pct"] is None else r["pct"],
            ]
            for r in report["rows"]
        ]
        grid_header = ["No. Absen", "Nama Lengkap"] + week_cols
        grid_data = [
            [r["roll_no"] or "", r["full_name"]] + [r["cells"].get(d, "") for d in report["week_days"]]
            for r in report["rows"]
        ]
        return (
            _xlsx_bytes([
                (f"Mingguan {start:%d-%m}", summary_header, summary_data),
                ("Per Hari", grid_header, grid_data),
            ]),
            XLSX_MIME,
            name + ".xlsx",
        )
    return _csv_bytes(header, data), CSV_MIME, name + ".csv"


def monthly_export(report, fmt):
    y, m = report["year"], report["month"]
    header = ["No. Absen", "Nama Lengkap", "Hadir", "Izin", "Sakit", "Alpa", "Belum tercatat", "% Hadir"]
    data = [[r["roll_no"] or "", r["full_name"], *[r["counts"][k] for k in STATUS_ORDER], r["belum"],
             "" if r["pct"] is None else r["pct"]] for r in report["rows"]]
    name = f"rekap-bulanan-{y}-{m:02d}"
    grid_header = ["No. Absen", "Nama Lengkap"] + [d[8:10] + "/" + d[5:7] for d in report["days"]]
    grid = [[r["roll_no"] or "", r["full_name"]] + [r["cells"].get(d, "") for d in report["days"]] for r in report["rows"]]
    if fmt == "xlsx":
        return (_xlsx_bytes([(f"Ringkasan {timeutil.MONTHS_ID[m-1]}", header, data), ("Per tanggal", grid_header, grid)]),
                XLSX_MIME, name + ".xlsx")
    return _csv_bytes(header, data), CSV_MIME, name + ".csv"


def student_export(student, history, fmt):
    header = ["Tanggal", "Hari", "Status", "Waktu Presensi", "Sumber", "Keterangan"]
    data = [[r["att_date"], timeutil.DAYS_ID[date.fromisoformat(r["att_date"]).weekday()],
             _status(r["status"]), timeutil.fmt_time(r["submitted_at"]),
             _source_label(r["source"]), r["note"] or ""] for r in history]
    name = "riwayat-" + "".join(c.lower() if c.isalnum() else "-" for c in student["full_name"]).strip("-")[:40]
    if fmt == "xlsx":
        return _xlsx_bytes([("Riwayat", header, data)]), XLSX_MIME, name + ".xlsx"
    return _csv_bytes(header, data), CSV_MIME, name + ".csv"
