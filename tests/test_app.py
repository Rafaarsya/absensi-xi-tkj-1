"""Uji otomatis aplikasi Absensi Kelas XI TKJ 1 (Siswa + Admin Kelas + Fitur Baru).
Jalankan: python -m unittest discover -s tests -t . -v"""
import io
import os
import re
import shutil
import sqlite3
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

from app import create_app, db as dbmod, security
from app.services import settings as settings_svc
from app.services import students as stu

WIB = timezone(timedelta(hours=7))
ADMIN_KEY = "kunci-admin-RAHASIA-4821"
ADMIN = ("Ketua Kelas XI TKJ 1", "Sandi-Admin-Uji-9")
STATIC_DIR = Path(__file__).resolve().parent.parent / "app" / "static"


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.clock = {"now": datetime(2026, 10, 26, 6, 45, tzinfo=WIB)}  # Senin 26 Okt 2026 06:45 WIB
        self.app = create_app({
            "APP_ENV": "testing",
            "TESTING": True,
            "SECRET_KEY": "x" * 40,
            "DATABASE": os.path.join(self.tmp, "t.db"),
            "ADMIN_ACCESS_KEY": ADMIN_KEY,
            "PASSWORD_HASH_METHOD": "pbkdf2:sha256:1000",
            "SESSION_COOKIE_SECURE": False,
            "NOW_FUNC": lambda: self.clock["now"],
        })
        conn = dbmod.connect(self.app.config["DATABASE"])
        dbmod.migrate(conn)
        self.class_id = dbmod.seed_class(conn, "XI TKJ 1", "SMKN 8 Kota Bekasi")
        conn.close()
        self.passwords = {}
        with self.app.app_context():
            db = dbmod.get_db()
            db.execute(
                "INSERT INTO users(role, full_name, name_key, password_hash, is_active, must_change_password, created_at) "
                "VALUES ('admin', ?, ?, ?, 1, 0, '2026-01-01T00:00:00+07:00')",
                (ADMIN[0], security.name_key(ADMIN[0]), security.hash_password(ADMIN[1])),
            )
            db.commit()
            for i, name in enumerate(["AATIF KHOIRUDDIN AMIN", "ALMIRA REGINA PUTRI", "BARAKA RAMADHANI"], start=1):
                sid, temp, err = stu.create_student(db, self.class_id, name, i)
                self.assertFalse(err)
                self.passwords[name] = (sid, temp)
        self.client = self.app.test_client()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def open_schedule(self, open_="06:30", close="07:15", days=range(7)):
        with self.app.app_context():
            settings_svc.save(
                dbmod.get_db(),
                {f"sched_{i}_{k}": v for i in days for k, v in (("open", open_), ("close", close))},
            )

    def token(self, client=None):
        client = client or self.client
        with client.session_transaction() as s:
            if "csrf" not in s:
                s["csrf"] = "t" * 43
            return s["csrf"]

    def post(self, url, data=None, client=None, **kw):
        client = client or self.client
        data = dict(data or {})
        data["csrf_token"] = self.token(client)
        return client.post(url, data=data, **kw)

    def login(self, role, name, password, key=None, client=None):
        client = client or self.client
        client.get(f"/login?role={role}")
        data = {"role": role, "full_name": name, "password": password}
        if key is not None:
            data["admin_key"] = key
        return self.post("/login", data, client=client)

    def login_admin(self, client=None):
        return self.login("admin", ADMIN[0], ADMIN[1], ADMIN_KEY, client)

    def login_student(self, name="AATIF KHOIRUDDIN AMIN", client=None, new_password="Sandi-Baru-77x"):
        client = client or self.client
        sid, temp = self.passwords[name]
        r = self.login("siswa", name, temp, client=client)
        self.assertEqual(r.status_code, 302)
        if new_password:
            r = self.post(
                "/akun/password",
                {"current_password": temp, "new_password": new_password, "confirm_password": new_password},
                client=client,
            )
            self.assertEqual(r.status_code, 302)
        return new_password

    def count(self, sql, args=()):
        with self.app.app_context():
            return dbmod.get_db().execute(sql, args).fetchone()[0]

    def new_client(self):
        return self.app.test_client()


class AuthTests(Base):
    def test_student_login_success_and_forced_password_change(self):
        sid, temp = self.passwords["AATIF KHOIRUDDIN AMIN"]
        r = self.login("siswa", "aatif  khoiruddin amin", temp)
        self.assertEqual(r.status_code, 302)
        r = self.client.get("/siswa/")
        self.assertEqual(r.status_code, 302)
        self.assertIn("/akun/password", r.headers["Location"])
        r = self.post("/akun/password", {"current_password": temp, "new_password": "12345678", "confirm_password": "12345678"})
        self.assertEqual(r.status_code, 422)
        r = self.post("/akun/password", {"current_password": temp, "new_password": "Sandi-Baru-77x", "confirm_password": "Sandi-Baru-77x"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.client.get("/siswa/").status_code, 200)

    def test_admin_login_requires_key_and_matches_role(self):
        self.assertEqual(self.login("admin", *ADMIN, key="kunci-salah").status_code, 401)
        self.assertEqual(self.login("admin", *ADMIN, key="", client=self.new_client()).status_code, 401)
        self.assertEqual(self.login_admin(self.new_client()).status_code, 302)
        # peran tidak boleh silang
        sid, temp = self.passwords["AATIF KHOIRUDDIN AMIN"]
        self.assertEqual(self.login("admin", "AATIF KHOIRUDDIN AMIN", temp, ADMIN_KEY, self.new_client()).status_code, 401)
        self.assertEqual(self.login("siswa", ADMIN[0], ADMIN[1], client=self.new_client()).status_code, 401)

    def test_login_page_texts_logo_and_dark_theme_toggle(self):
        """Memastikan logo kelas, tombol switch tema gelap, dan teks halaman login admin sesuai spesifikasi."""
        admin_login = self.client.get("/login?role=admin").get_data(as_text=True)
        self.assertIn("logo-kelas.png", admin_login)
        self.assertIn("data-theme-toggle", admin_login)
        self.assertIn("Khusus pengurus yang ditunjuk (Ketua Kelas dan Sekretaris)", admin_login)
        self.assertIn("Masukkan akun admin yang terdaftar", admin_login)
        self.assertNotIn("h-admin_key", admin_login)
        self.assertIn("Lupa kata sandi atau kunci akses? Hubungi admin web.", admin_login)
        self.assertTrue((STATIC_DIR / "img" / "logo-kelas.png").is_file())

    def test_admin_key_never_exposed_to_client(self):
        bodies = [
            self.client.get("/login?role=admin").get_data(as_text=True),
            self.client.get("/login?role=siswa").get_data(as_text=True),
        ]
        for f in STATIC_DIR.rglob("*"):
            if f.is_file() and f.suffix in (".css", ".js", ".svg"):
                bodies.append(f.read_text(encoding="utf-8", errors="ignore"))
        self.login_admin()
        for path in ("/admin/", "/admin/presensi", "/admin/siswa", "/admin/laporan/mingguan", "/admin/pengaturan"):
            bodies.append(self.client.get(path).get_data(as_text=True))
        for b in bodies:
            self.assertNotIn(ADMIN_KEY, b)

    def test_login_rate_limit_and_session_timeout(self):
        sid, temp = self.passwords["AATIF KHOIRUDDIN AMIN"]
        for _ in range(5):
            self.assertEqual(self.login("siswa", "AATIF KHOIRUDDIN AMIN", "salah-salah-1").status_code, 401)
        self.assertEqual(self.login("siswa", "AATIF KHOIRUDDIN AMIN", temp).status_code, 429)
        self.clock["now"] += timedelta(minutes=16)
        self.assertEqual(self.login("siswa", "AATIF KHOIRUDDIN AMIN", temp).status_code, 302)


class CspRegressionTests(unittest.TestCase):
    def test_templates_have_no_inline_style_or_handlers(self):
        tpl = Path(__file__).resolve().parent.parent / "app" / "templates"
        for f in tpl.rglob("*.html"):
            text = f.read_text(encoding="utf-8")
            self.assertNotRegex(text, r"\sstyle\s*=", f.name)
            self.assertNotRegex(text, r"\son[a-z]+\s*=\s*[\"']", f.name)
            self.assertNotRegex(text, r"<script(?![^>]*\bsrc=)", f.name)
            self.assertNotRegex(text, r"<style", f.name)
            self.assertNotRegex(text, r"https?://(?!www\.w3\.org)", f.name)


class StudentPortalTests(Base):
    def test_student_cannot_open_admin_pages(self):
        self.login_student()
        for p in ("/admin/", "/admin/presensi", "/admin/siswa", "/admin/laporan/mingguan", "/admin/pengaturan"):
            self.assertEqual(self.client.get(p).status_code, 403, p)
        for p, data in [
            ("/admin/presensi/tandai-semua-hadir", {}),
            ("/admin/presensi/ubah-cepat", {"siswa": 2, "status": "H"}),
            ("/admin/siswa/2/reset-password", {}),
        ]:
            self.assertEqual(self.post(p, data).status_code, 403, p)

    def test_student_attendance_and_geofence(self):
        self.open_schedule()
        with self.app.app_context():
            settings_svc.save(dbmod.get_db(), {"geo_enabled": "1", "geo_lat": "-6.2", "geo_lng": "106.8", "geo_radius": "100"})
        self.login_student()
        self.assertEqual(self.post("/siswa/presensi", {"status": "H"}).status_code, 422)  # tanpa GPS
        self.assertEqual(self.post("/siswa/presensi", {"status": "H", "lat": "-6.2003", "lng": "106.8"}).status_code, 302)
        self.assertEqual(self.count("SELECT COUNT(*) FROM attendance_records WHERE status='H'"), 1)
        # Kirim kedua kali ditolak/diabaikan
        self.assertEqual(self.post("/siswa/presensi", {"status": "I", "note": "coba ubah"}).status_code, 302)
        self.assertEqual(self.count("SELECT status FROM attendance_records"), "H")


class AdminPortalTests(Base):
    ADMIN_PATHS = [
        "/admin/",
        "/admin/presensi",
        "/admin/siswa",
        "/admin/siswa/baru",
        "/admin/siswa/2",
        "/admin/presensi/koreksi?siswa=2",
        "/admin/laporan/harian",
        "/admin/laporan/mingguan",
        "/admin/laporan/bulanan",
        "/admin/pengaturan",
    ]

    def test_all_admin_pages_render(self):
        self.open_schedule()
        self.login_admin()
        for p in self.ADMIN_PATHS + ["/admin/laporan", "/admin/siswa?q=aatif&status=aktif", "/admin/presensi?status=belum"]:
            r = self.client.get(p, follow_redirects=True)
            self.assertEqual(r.status_code, 200, p)

    def test_mark_all_present_then_edit_absent_students(self):
        """Alur cepat Admin: Klik 'Tandai Semua Hadir', lalu ubah status siswa yang tidak masuk."""
        self.login_admin()
        r = self.post("/admin/presensi/tandai-semua-hadir", {"tanggal": "2026-10-26"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.count("SELECT COUNT(*) FROM attendance_records WHERE status='H' AND att_date='2026-10-26'"), 3)

        almira_id = self.passwords["ALMIRA REGINA PUTRI"][0]
        baraka_id = self.passwords["BARAKA RAMADHANI"][0]

        self.assertEqual(self.post("/admin/presensi/ubah-cepat", {
            "siswa": almira_id, "tanggal": "2026-10-26", "status": "S", "note": "Demam tinggi"
        }).status_code, 302)
        self.assertEqual(self.post("/admin/presensi/ubah-cepat", {
            "siswa": baraka_id, "tanggal": "2026-10-26", "status": "I", "note": "Acara keluarga"
        }).status_code, 302)

        self.assertEqual(self.count("SELECT COUNT(*) FROM attendance_records WHERE status='H'"), 1)
        self.assertEqual(self.count("SELECT COUNT(*) FROM attendance_records WHERE status='S'"), 1)
        self.assertEqual(self.count("SELECT COUNT(*) FROM attendance_records WHERE status='I'"), 1)
        self.assertEqual(self.count("SELECT COUNT(*) FROM attendance_corrections"), 2)

    def test_weekly_report_and_exports(self):
        self.login_admin()
        aatif_id = self.passwords["AATIF KHOIRUDDIN AMIN"][0]
        almira_id = self.passwords["ALMIRA REGINA PUTRI"][0]

        # Senin 26 Okt 2026
        self.post("/admin/presensi/tandai-semua-hadir", {"tanggal": "2026-10-26"})
        self.post("/admin/presensi/ubah-cepat", {"siswa": almira_id, "tanggal": "2026-10-26", "status": "S", "note": "Flu"})

        # Selasa 27 Okt 2026
        self.clock["now"] = datetime(2026, 10, 27, 7, 0, tzinfo=WIB)
        self.login_admin()
        self.post("/admin/presensi/tandai-semua-hadir", {"tanggal": "2026-10-27"})
        self.post("/admin/presensi/ubah-cepat", {"siswa": aatif_id, "tanggal": "2026-10-27", "status": "I", "note": "Lomba"})

        html = self.client.get("/admin/laporan/mingguan?minggu=2026-10-27").get_data(as_text=True)
        self.assertIn("26 Okt – 1 Nov 2026", html)
        self.assertRegex(html, r"Hari efektif</dt><dd>2</dd>")

        csv_resp = self.client.get("/admin/laporan/mingguan?minggu=2026-10-26&unduh=csv")
        self.assertEqual(csv_resp.status_code, 200)
        csv_text = csv_resp.get_data(as_text=True)
        self.assertIn("AATIF KHOIRUDDIN AMIN,H,I,,,,1,1,0,0,0,50.0", csv_text)
        self.assertIn("ALMIRA REGINA PUTRI,S,H,,,,1,0,1,0,0,50.0", csv_text)
        self.assertIn("BARAKA RAMADHANI,H,H,,,,2,0,0,0,0,100.0", csv_text)

        xlsx_resp = self.client.get("/admin/laporan/mingguan?minggu=2026-10-26&unduh=xlsx")
        self.assertEqual(xlsx_resp.status_code, 200)
        wb = load_workbook(io.BytesIO(xlsx_resp.data))
        self.assertEqual(len(wb.sheetnames), 2)

    def test_student_account_creation_and_password_reset(self):
        self.login_admin()
        r = self.post("/admin/siswa/baru", {"full_name": "SITI NURHALIZA", "roll_no": "", "nis": "12345"})
        self.assertEqual(r.status_code, 200)
        temp = re.search(r'class="credential">([^<]+)<', r.get_data(as_text=True)).group(1)
        self.assertEqual(len(temp), 10)
        c = self.new_client()
        self.assertEqual(self.login("siswa", "siti nurhaliza", temp, client=c).status_code, 302)
        self.assertIn("/akun/password", c.get("/siswa/").headers["Location"])


if __name__ == "__main__":
    unittest.main()
